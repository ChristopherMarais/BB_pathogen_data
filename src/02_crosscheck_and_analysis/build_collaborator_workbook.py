"""Streamlined 2-Sheet Candidate Pathogen Review Workbook for Collaborators:
  - Sheet 1 ('100 Random sample'): 100 randomly sampled focused candidate records for collaborator precision verification (yellow review columns with interactive drop-downs).
  - Sheet 2 ('All candidate records'): Complete catalog of all 1,665 candidate records from the automated search and database mining, with filterable subset/group, beetle taxonomy, locators, and supporting quotes.

Usage:
    python build_collaborator_workbook.py <candidate_table.csv> <out.xlsx>
"""
from pathlib import Path
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SEED, N_SAMPLE = 20261026, 100
src = Path(sys.argv[1]).resolve()
out = Path(sys.argv[2]).resolve()

def clean_df(df):
    df = df.copy()
    for c in df.columns:
        if df[c].dtype == object or str(df[c].dtype).startswith("str"):
            df[c] = df[c].fillna("").astype(str).str.replace(ILLEGAL_CHARACTERS_RE, " ", regex=True)
    return df

T = clean_df(pd.read_csv(src, dtype=str).fillna(""))
B = lambda s: s.astype(str).str.lower().isin(["true", "1", "yes"])
T["focused_b"], T["new_pair_b"], T["new_beetle_b"] = B(T.focused), B(T.new_pair), B(T.new_beetle_species)
lit = T.stratum.str.contains("literature", case=False)

# Stratify groups for filtering
T["sheet_temp"] = np.where(~lit | ~T.focused_b, "5", "4")
T.loc[lit & T.beetle_subfamily.eq("Platypodinae"), "sheet_temp"] = "2"
T.loc[lit & T.focused_b & T.beetle.eq("Hypothenemus hampei"), "sheet_temp"] = "3"
pool = T[lit & T.focused_b & T.sheet_temp.isin(["3", "4"])]
sample_ids = pool.sample(n=min(N_SAMPLE, len(pool)), random_state=SEED).index

def assign_group(row):
    if row["beetle_subfamily"] == "Platypodinae":
        return "Platypodinae"
    elif row["beetle"] == "Hypothenemus hampei":
        return "Coffee berry borer (Hypothenemus hampei)"
    elif row["focused_b"]:
        return "Other focused literature candidates"
    else:
        return "Database & broader literature"

T["in_sample"] = np.where(T.index.isin(sample_ids), "Yes", "")
T["group_label"] = [assign_group(r) for _, r in T.iterrows()]
T["link"] = np.where(T.doi.ne(""), "https://doi.org/" + T.doi.str.replace("https://doi.org/", "", regex=False), T.url)
T["stratum"] = T["stratum"].str.replace("literature read by the KG", "automated literature reading", regex=False)
T["new_beetle"] = np.where(T.new_beetle_b, "yes", "")
T["new_pair"] = np.where(T.new_pair_b, "yes", "")
T["relation_read"] = T.relation_read.where(T.relation_read.ne(""), T.relation)

COLS = [
    ("Record ID", "fact_id", 16),
    ("In 100 Sample?", "in_sample", 14),
    ("Subset / Group", "group_label", 26),
    ("Source Stratum", "stratum", 22),
    ("Beetle", "beetle", 24),
    ("Subfamily", "beetle_subfamily", 14),
    ("Organism as written", "pathogen_as_written", 26),
    ("Organism (GBIF name)", "pathogen_resolved_gbif", 26),
    ("Category", "category", 12),
    ("Relation as read", "relation_read", 14),
    ("Year", "year", 7),
    ("Publication", "citation", 60),
    ("DOI / link", "link", 30),
    ("Page / Locator", "locator", 16),
    ("Source ID", "source_id", 24),
    ("Language", "language", 9),
    ("Beetle new to dataset", "new_beetle", 10),
    ("Pair new to dataset", "new_pair", 10),
    ("Supporting sentence (lab only)", "supporting_quote_INTERNAL_ONLY", 70),
    ("MEETS INCLUSION RULES?", None, 14),
    ("Beetle correct?", None, 10),
    ("Organism correct?", None, 10),
    ("If not included: reason", None, 26),
    ("Country / region (if stated)", None, 18),
    ("Notes", None, 30),
]

YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="D9E1F2")
reasons = ('"not called pathogen/parasite,associate only (no harm),cites another work,review/secondary source,'
           'wrong beetle,wrong organism,not one of the six groups,already in dataset,other"')

wb = Workbook()
# Remove default sheet
wb.remove(wb.active)

sheets_to_create = [
    ("100 Random sample", T.loc[sample_ids].sort_values(["group_label", "citation", "beetle"])),
    ("All candidate records", T.sort_values(["group_label", "citation", "beetle"])),
]

for title, D in sheets_to_create:
    s = wb.create_sheet(title)
    for j, (h, _, w) in enumerate(COLS, 1):
        c = s.cell(row=1, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = YELLOW if COLS[j - 1][1] is None else HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
        s.column_dimensions[get_column_letter(j)].width = w

    for i, (_, r) in enumerate(D.iterrows(), 2):
        for j, (_, col, _) in enumerate(COLS, 1):
            if col is None:
                s.cell(row=i, column=j).fill = YELLOW
                continue
            v = r[col] if col in r else ""
            c = s.cell(row=i, column=j, value=int(v) if col == "year" and str(v).isdigit() else v)
            c.alignment = Alignment(wrap_text=col in ("citation", "supporting_quote_INTERNAL_ONLY"), vertical="top")
            if col == "link" and str(v).startswith("http"):
                c.hyperlink = v
                c.font = Font(color="0563C1", underline="single")
            if col in ("beetle", "pathogen_as_written", "pathogen_resolved_gbif"):
                c.font = Font(italic=True)

    n = len(D) + 1
    # Reviewer columns: T (20), U (21), V (22), W (23)
    for col_letter, formula in (("T", '"Yes,No,Unsure"'), ("U", '"Yes,No"'), ("V", '"Yes,No"'), ("W", reasons)):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        s.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{max(n, 2)}")
    s.freeze_panes = "F2"
    s.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{max(n, 1)}"

out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Successfully generated 2-sheet workbook: {out} ({len(wb.sheetnames)} sheets: {wb.sheetnames})")
